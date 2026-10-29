"""Read the BALLOP stock-taking workbook (sheets STOK TERKINI + DATA SO + SUMMARY) into one JSON file.
usage: python tools/stocktaking-to-json.py "STOCK TAKING - BALLOP.xlsx" out.json
Needs: pip install openpyxl"""
import json
import re
import sys
from datetime import datetime

import openpyxl

src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
norm = lambda s: re.sub(r"\s+", " ", str(s or "")).strip()


def size_txt(v):
    if v is None or v == "":
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace(",", ".").strip()


# 1. warehouse SKU codes per item name, from the sales-order export
codes = {}
ws = wb["DATA SO"]
head = None
for row in ws.iter_rows(values_only=True):
    if head is None:
        if row and "Item SKU" in row:
            head = list(row)
        continue
    name, sku = row[head.index("Item Name")], row[head.index("Item SKU")]
    if name and sku:
        codes[norm(name).upper()] = norm(sku)

# 2. stock date = "Data SO s/d" on SUMMARY (fallback: today)
stock_date = datetime.now().strftime("%Y-%m-%d")
for row in wb["SUMMARY"].iter_rows(values_only=True, max_row=6):
    cells = list(row)
    for i, c in enumerate(cells):
        if isinstance(c, str) and "Data SO" in c and i + 1 < len(cells) and isinstance(cells[i + 1], datetime):
            stock_date = cells[i + 1].strftime("%Y-%m-%d")

# 3. current stock per SKU
items = []
ws = wb["STOK TERKINI"]
head = None
for row in ws.iter_rows(values_only=True):
    if head is None:
        if row and row[0] == "NO":
            head = [norm(c) for c in row]
        continue
    if not row or not row[1] or not row[2]:
        continue
    key = norm(row[1])
    items.append({
        "key": key,
        "product": norm(row[2]).upper(),
        "color": norm(row[3]).upper(),
        "sizeKr": size_txt(row[4]),
        "sizeEu": size_txt(row[5]),
        "good": int(row[6] or 0),
        "damaged": int(row[7] or 0),
        "sku": codes.get(key.upper(), ""),
    })

json.dump({"date": stock_date, "warehouse": "GBI Bintaro", "items": items}, open(dst, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print(f"{len(items)} SKU, {sum(1 for i in items if i['sku'])} dengan kode gudang, tanggal {stock_date} -> {dst}")
