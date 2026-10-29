"""Turn the BALLOP operations workbook + stock-taking workbook into admin data (one JSON bundle).
usage: python tools/ops-to-json.py OPERASIONAL.xlsx STOCKTAKING.xlsx out.json
Needs: pip install openpyxl

  stock taking : STOK TERKINI (stock per SKU) + DATA SO (warehouse SKU codes)       -> products, stock3pl
  operasional  : 83/84/85 RAW (order lines Shopee / TikTok / Tokopedia)            -> orders, settlements
                 10 Omzet Online, "Closing bulan" rows for months without RAW data  -> orders (monthly recap)
                 13 Iklan -> ads, 14 Kreator -> kol, 20 Omzet Offline -> offline,
                 22 Biaya Offline -> expenses, 26 Kas -> cashflow
"""
import json
import re
import sys
from collections import defaultdict
from datetime import datetime

import openpyxl

ops_path, st_path, out_path = sys.argv[1], sys.argv[2], sys.argv[3]
ops = openpyxl.load_workbook(ops_path, read_only=True, data_only=True)
stw = openpyxl.load_workbook(st_path, read_only=True, data_only=True)

norm = lambda s: re.sub(r"\s+", " ", str(s if s is not None else "")).strip()
idn = lambda s: norm(str(s).split("/")[-1]) if "/" in str(s or "") else norm(s)  # "한국어 / Indonesia" -> Indonesia


def num(v):
    if v is None or v == "" or v == "-":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.,-]", "", str(v))
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return float(s.replace(",", "."))
    except ValueError:
        return 0.0


def day(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    m = re.match(r"(\d{4})-(\d{2})(?:-(\d{2}))?", str(v or ""))
    return f"{m.group(1)}-{m.group(2)}-{m.group(3) or '01'}" if m else ""


def size_txt(v):
    if v in (None, ""):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace(",", ".").strip()


def rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def table(ws, first_header):
    """rows after the header row whose first cell is `first_header`"""
    out, on = [], False
    for r in ws.iter_rows(values_only=True):
        if not on:
            on = r and norm(r[0]) == first_header
            continue
        if any(c not in (None, "") for c in r):
            out.append(list(r) + [None] * 30)
    return out


# ---------------- stock taking -> products + stock3pl ----------------
codes = {}
head = None
for r in stw["DATA SO"].iter_rows(values_only=True):
    if head is None:
        head = list(r) if r and "Item SKU" in r else None
        continue
    if r[head.index("Item Name")] and r[head.index("Item SKU")]:
        codes[norm(r[head.index("Item Name")]).upper()] = norm(r[head.index("Item SKU")])

stock_date = datetime.now().strftime("%Y-%m-%d")
for r in stw["SUMMARY"].iter_rows(values_only=True, max_row=6):
    for i, c in enumerate(r):
        if isinstance(c, str) and "Data SO" in c and i + 1 < len(r) and isinstance(r[i + 1], datetime):
            stock_date = r[i + 1].strftime("%Y-%m-%d")

slug = lambda s: re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")
products, stock3pl = {}, []
for r in table(stw["STOK TERKINI"], "NO"):
    if not r[1] or not r[2]:
        continue
    key, prod, color = norm(r[1]), norm(r[2]).upper(), norm(r[3]).upper()
    pid = slug(f"{prod} {color}")
    eu, kr = size_txt(r[5]), size_txt(r[4])
    p = products.setdefault(pid, {"id": pid, "name": prod.replace("BALLOP ", "", 1), "series": color, "variants": [], "sizes": []})
    p["variants"].append({"size": eu, "sizeKr": kr, "sku": codes.get(key.upper(), "")})
    if eu and eu not in p["sizes"]:
        p["sizes"].append(eu)
    stock3pl.append({
        "date": stock_date, "warehouse": "GBI Bintaro", "productId": pid, "sku": codes.get(key.upper(), ""),
        "name": f"{prod} {color}", "size": eu, "available": int(num(r[6])), "reserved": 0, "damaged": int(num(r[7])), "note": key,
    })

# ---------------- RAW marketplace order lines ----------------
RAW = {
    "shopee": ("83_RAW_Shopee", {"payout": 15, "fees": 14, "note": 16}),
    "tiktok": ("84_RAW_Tiktok", {"payout": 17, "fees": 16, "note": 18, "after": 7, "before": 8, "sellerDisc": 9, "platDisc": 10}),
    "tokopedia": ("85_RAW_Tokopedia", {"payout": 17, "fees": 16, "note": 18, "after": 7, "before": 8, "sellerDisc": 9, "platDisc": 10}),
}
orders, settlements = [], []
raw_months = defaultdict(set)
for ch, (sheet, col) in RAW.items():
    prev_oid, note, d = None, "", ""
    line_no = defaultdict(int)
    for r in rows(ops[sheet]):
        r = r + [None] * 30
        oid = norm(r[24]) or norm(r[1])
        date = next((x for x in (r[23], r[0], r[25]) if isinstance(x, datetime)), None)
        qty = int(num(r[5]))
        if not oid or not r[2] or qty <= 0 or (not date and oid != prev_oid):
            continue  # empty template rows at the bottom of the sheet
        first = oid != prev_oid
        if first:  # a new order starts: its note, date and money sit on this row
            prev_oid, note, d = oid, norm(r[col["note"]]).lower() or "sales", date.strftime("%Y-%m-%d")
        if note not in ("sales", "refund"):
            continue  # free samples / add-on gifts are not sales
        state = "returned" if note == "refund" else "completed"
        unit = num(r[6])
        if not first and settlements and settlements[-1]["orderId"] == oid:
            # multi-item Shopee orders carry income + fees on every item row
            settlements[-1]["payout"] += num(r[col["payout"]])
            settlements[-1]["fees"] -= abs(num(r[col["fees"]]))
        if first:
            payout, fees = num(r[col["payout"]]), num(r[col["fees"]])
            settlements.append({"key": f"{ch}:{oid}", "channel": ch, "orderId": oid, "type": note.title(), "date": d, "releaseDate": "",
                                "productPrice": num(r[col["after"]]) if "after" in col else 0, "payout": payout, "fees": -abs(fees)})
        line_no[oid] += 1
        net = unit * qty
        orders.append({
            "key": f"{ch}:{oid}:{line_no[oid]}", "channel": ch, "orderId": oid, "status": note.title(), "state": state, "date": d,
            "sku": "", "productName": norm(r[2]), "variation": f"{norm(r[3])}, {size_txt(r[4])}".strip(", "), "size": size_txt(r[4]),
            "qty": qty, "unitPrice": unit, "gross": net, "net": net,
            "sellerDiscount": num(r[col["sellerDisc"]]) if first and "sellerDisc" in col else 0,
            "platformDiscount": num(r[col["platDisc"]]) if first and "platDisc" in col else 0,
        })
        raw_months[ch].add(d[:7])

# monthly recap rows for months the RAW sheets do not cover
CH = {"tiktok shop": "tiktok", "shopee": "shopee", "tokopedia": "tokopedia", "lazada": "lazada", "blibli": "blibli"}
for r in table(ops["10_IN_온라인매출_Omzet Online"], "No"):
    if "월마감" not in norm(r[1]) and "Closing" not in norm(r[1]):
        continue
    ch = CH.get(norm(r[5]).lower())
    month = day(r[2])[:7]
    gmv, n_orders = num(r[7]), int(num(r[8]))
    if not ch or not month or gmv <= 0 or month in raw_months[ch]:
        continue
    orders.append({"key": f"{ch}:rekap:{month}", "channel": ch, "orderId": f"REKAP-{month}", "status": "Rekap bulanan", "state": "completed",
                   "date": f"{month}-01", "sku": "", "productName": "(rekap bulanan, semua produk)", "variation": "", "size": "",
                   "qty": n_orders, "unitPrice": 0, "gross": gmv, "net": gmv, "note": "rekap 10_Omzet Online"})
    settlements.append({"key": f"{ch}:rekap:{month}", "channel": ch, "orderId": f"REKAP-{month}", "type": "Rekap bulanan", "date": f"{month}-01",
                        "releaseDate": "", "productPrice": gmv, "payout": num(r[13]), "fees": -abs(num(r[12]))})

# ---------------- ads ----------------
ADCH = {"tiktok shop": "tiktok", "tiktok": "tiktok", "shopee": "shopee", "lazada": "lazada", "blibli": "blibli", "meta": "meta", "facebook": "meta", "instagram": "meta", "google": "google"}
ads = []
for r in table(ops["13_IN_광고라인_Iklan"], "No"):
    if not r[6] and not r[8]:
        continue
    ch = next((v for k, v in ADCH.items() if k in norm(r[2]).lower()), "other")
    ads.append({"date": day(r[7]) or f"{norm(r[1])}-01", "channel": ch, "campaign": norm(r[6]) or norm(r[4]),
                "spend": num(r[8]), "revenue": num(r[9]), "note": f"{idn(r[3])} · {idn(r[4])} · {idn(r[12])}"[:200]})

# ---------------- creators / KOL ----------------
kol = []
for r in table(ops["14_IN_크리에이터_Kreator"], "No"):
    if not r[2]:
        continue
    plat = norm(r[4]).lower()
    kol.append({
        "date": f"{norm(r[1])}-01" if re.match(r"\d{4}-\d{2}$", norm(r[1])) else day(r[1]), "name": norm(r[2]), "handle": norm(r[2]),
        "platform": "tiktok" if "tik" in plat else "instagram" if "insta" in plat else "youtube" if "you" in plat else "other",
        "type": idn(r[3]), "gmv": num(r[5]), "orders": int(num(r[6])), "sold": int(num(r[7])),
        "cost": num(r[9]) + num(r[10]) + num(r[11]), "qty": 1, "shipStatus": "diterima", "productId": f"n:{norm(r[17])} {norm(r[18])}".strip() if norm(r[17]) else "",
        "contentStatus": "posted" if num(r[5]) > 0 else "menunggu",
        "note": " · ".join(x for x in [norm(r[17]), norm(r[18]), size_txt(r[19]), norm(r[16])] if x)[:300],
    })

# ---------------- offline sales ----------------
def off_type(sale_type):
    s = sale_type.lower()
    return "event" if ("event" in s or "이벤트" in s) else "distributor" if ("b2b" in s or "사입" in s) else "toko"

offline = []
for r in table(ops["20_IN_오프라인매출_Omzet Offline"], "No"):
    if not r[3] or not num(r[8]):
        continue
    offline.append({"date": day(r[1]), "outlet": norm(r[3]), "type": off_type(norm(r[5])), "qty": int(num(r[8])), "price": num(r[9]),
                    "discount": num(r[12]), "payment": "lunas", "status": "completed", "productId": f"n:{norm(r[7])}",
                    "note": " · ".join(x for x in [norm(r[7]), idn(r[5]), norm(r[17])] if x)[:200]})

# ---------------- offline costs ----------------
expenses = []
for r in table(ops["22_IN_오프라인비용_Biaya Offline"], "No"):
    cash, hpp = num(r[11]), num(r[12])
    if not (cash or hpp):
        continue
    expenses.append({"date": day(r[1]), "category": idn(r[3])[:40] or "Lainnya", "amount": cash or hpp,
                     "note": " · ".join(x for x in [norm(r[4]), norm(r[5]), norm(r[6]), "basis HPP" if not cash else "", norm(r[14])] if x)[:200]})

# ---------------- cash book ----------------
cashflow = []
for r in table(ops["26_IN_입출금_Kas"], "No"):
    if not isinstance(r[1], datetime) or not r[3]:
        continue
    cat = idn(r[5])
    if "saldo awal" in cat.lower():
        continue  # opening balances are not cash movements
    cin, cout = num(r[7]), num(r[8])
    if not (cin or cout):
        continue
    cashflow.append({"date": day(r[1]), "type": "in" if cin else "out", "category": cat[:40] or "Lainnya", "account": norm(r[3])[:40],
                     "amount": cin or cout, "ref": "", "note": " · ".join(x for x in [norm(r[4]), idn(r[6]), norm(r[12])] if x)[:200]})

bundle = {
    "products": list(products.values()), "stock3pl": stock3pl, "orders": orders, "settlements": settlements,
    "ads": ads, "kol": kol, "offline": offline, "expenses": expenses, "cashflow": cashflow,
}
json.dump(bundle, open(out_path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
for k, v in bundle.items():
    print(f"{k:12} {len(v)}")
print("raw months:", {k: sorted(v) for k, v in raw_months.items()})
